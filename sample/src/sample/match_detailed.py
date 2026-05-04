"""Sample 2: 1試合 = 1スライドの詳細版。

simple 版にスタジアム・観客数・得点者リスト・節（matchday）を加えた、
実用に近いレイアウト。python-pptx で図形・テキスト・色をどう積み上げるかを確認する。

実行:
    rye run python src/sample/match_detailed.py
出力:
    match_detailed.pptx, match_detailed.pdf
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from pdf_export import convert_to_pdf


MATCH = {
    "matchday": 35,
    "home": "Liverpool",
    "away": "Manchester City",
    "home_score": 3,
    "away_score": 2,
    "date": "2026-05-03 (Sun) 17:30",
    "stadium": "Anfield, Liverpool",
    "attendance": 53_394,
    "home_scorers": [
        ("M. Salah", 12),
        ("D. Núñez", 41),
        ("L. Díaz", 78),
    ],
    "away_scorers": [
        ("E. Haaland", 23),
        ("P. Foden", 65),
    ],
}

TEAM_COLORS = {
    "Arsenal": "EF0107",
    "Chelsea": "034694",
    "Liverpool": "C8102E",
    "Manchester City": "6CABDD",
    "Manchester United": "DA291C",
    "Tottenham": "132257",
    "Newcastle": "241F20",
    "Aston Villa": "670E36",
}

PL_PURPLE = "37003C"


def add_text(slide, left, top, width, height, text, *,
             size=18, bold=False, color="000000", align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def add_filled_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def add_team_panel(slide, left, top, width, height, team):
    color = TEAM_COLORS.get(team, "555555")
    panel = add_filled_rect(slide, left, top, width, height, color)
    tf = panel.text_frame
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = team
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("FFFFFF")


def add_scorer_list(slide, left, top, width, height, scorers, *, align=PP_ALIGN.LEFT):
    """得点者を「名前 12'」形式の箇条書きで描画."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (name, minute) in enumerate(scorers):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = f"⚽ {name}  {minute}'"
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor.from_string("222222")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ヘッダー帯（プレミアリーグカラー）
    add_filled_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.9), PL_PURPLE)
    add_text(slide, Inches(0), Inches(0.2), Inches(13.33), Inches(0.6),
             f"Premier League · Matchday {MATCH['matchday']}",
             size=24, bold=True, color="FFFFFF")

    # チームパネル + スコア
    add_team_panel(slide, Inches(0.5), Inches(1.3), Inches(4.5), Inches(2.0), MATCH["home"])
    add_team_panel(slide, Inches(8.33), Inches(1.3), Inches(4.5), Inches(2.0), MATCH["away"])

    score = f"{MATCH['home_score']}  -  {MATCH['away_score']}"
    add_text(slide, Inches(5.0), Inches(1.4), Inches(3.33), Inches(1.4),
             score, size=72, bold=True, color="111111")
    add_text(slide, Inches(5.0), Inches(2.7), Inches(3.33), Inches(0.5),
             "FT", size=18, bold=True, color="888888")

    # メタ情報（日付・スタジアム・観客数）
    add_text(slide, Inches(0.5), Inches(3.5), Inches(12.33), Inches(0.4),
             MATCH["date"], size=16, color="333333")
    add_text(slide, Inches(0.5), Inches(3.95), Inches(12.33), Inches(0.4),
             f"{MATCH['stadium']} · Att. {MATCH['attendance']:,}",
             size=14, color="666666")

    # 得点者ヘッダ
    add_filled_rect(slide, Inches(0.5), Inches(4.6), Inches(12.33), Inches(0.5), "EEEEEE")
    add_text(slide, Inches(0.5), Inches(4.65), Inches(12.33), Inches(0.4),
             "Goals", size=16, bold=True, color="333333")

    # 得点者リスト（左:ホーム / 右:アウェイ）
    add_scorer_list(slide, Inches(0.7), Inches(5.2), Inches(6.0), Inches(2.0),
                    MATCH["home_scorers"], align=PP_ALIGN.LEFT)
    add_scorer_list(slide, Inches(7.0), Inches(5.2), Inches(6.0), Inches(2.0),
                    MATCH["away_scorers"], align=PP_ALIGN.RIGHT)

    out = "match_detailed.pptx"
    prs.save(out)
    print(f"saved: {out}")
    pdf = convert_to_pdf(out)
    print(f"saved: {pdf.name}")


if __name__ == "__main__":
    build()
