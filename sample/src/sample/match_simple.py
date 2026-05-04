"""Sample 1: 1試合 = 1スライドの最小構成。

中央に「ホーム名 スコア アウェイ名」、上にリーグ名、下に日付だけを置く
シンプルなレイアウト。python-pptx の基本 API を確認するための入門サンプル。

実行:
    rye run python src/sample/match_simple.py
出力:
    match_simple.pptx
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


# 試合データ（ダミー）
MATCH = {
    "home": "Arsenal",
    "away": "Chelsea",
    "home_score": 2,
    "away_score": 1,
    "date": "2026-05-03 (Sun) 16:30",
}

# プレミアリーグの主なチームカラー（RGB hex）
TEAM_COLORS = {
    "Arsenal": "EF0107",
    "Chelsea": "034694",
    "Liverpool": "C8102E",
    "Manchester City": "6CABDD",
    "Manchester United": "DA291C",
    "Tottenham": "132257",
    "Newcastle": "241F20",
    "Aston Villa": "670E36",
}


def add_text(slide, left, top, width, height, text, *, size=18, bold=False, color="000000", align=PP_ALIGN.CENTER):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def add_team_block(slide, left, top, width, height, team):
    """チーム名をチームカラーの帯で表示する."""
    color = TEAM_COLORS.get(team, "555555")
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = team
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("FFFFFF")


def build():
    prs = Presentation()
    # 16:9 のワイド
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]  # 完全に空白のレイアウト
    slide = prs.slides.add_slide(blank)

    # 上部: リーグ名
    add_text(slide, Inches(0), Inches(0.4), Inches(13.33), Inches(0.6),
             "Premier League", size=24, bold=True, color="37003C")

    # 中央: チーム + スコア + チーム
    block_w, block_h = Inches(4.5), Inches(2.5)
    block_top = Inches(2.5)

    add_team_block(slide, Inches(0.5), block_top, block_w, block_h, MATCH["home"])
    add_team_block(slide, Inches(8.33), block_top, block_w, block_h, MATCH["away"])

    score_text = f"{MATCH['home_score']}  -  {MATCH['away_score']}"
    add_text(slide, Inches(5.0), Inches(2.9), Inches(3.33), Inches(1.8),
             score_text, size=72, bold=True, color="111111")

    # 下部: 日付
    add_text(slide, Inches(0), Inches(6.4), Inches(13.33), Inches(0.6),
             MATCH["date"], size=16, color="555555")

    out = "match_simple.pptx"
    prs.save(out)
    print(f"saved: {out}")


if __name__ == "__main__":
    build()
